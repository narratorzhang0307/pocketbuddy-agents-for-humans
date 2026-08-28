#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Pocket_Earth_阿里通义Arm决赛_完整产品文档与30页PPT脚本_2026-08-13.md"
OUTPUT = ROOT / "Pocket_Earth_阿里通义Arm决赛_完整产品文档与30页PPT脚本_2026-08-13.docx"

BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
GREEN = "00E676"
BLACK = "111111"
GRAY = "666666"
LIGHT = "F4F6F9"
CREAM = "FBF6E8"
PALE_GREEN = "E9F9F0"
GOLD = "C8A33A"
FONT_BODY = "Arial Unicode MS"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for name, attrs in edges.items():
        tag = "w:" + name
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn("w:" + key), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, name=FONT_BODY) -> None:
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), FONT_BODY)
    rfonts.set(qn("w:cs"), FONT_BODY)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str, default_size=10.5, default_color=BLACK) -> None:
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, default_size, True, default_color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, default_size - 0.5, False, DEEP_BLUE, FONT_BODY)
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F5")
            rpr.append(shd)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, default_size, False, default_color)


def set_paragraph_body(paragraph, after=8, line=1.333, justify=True) -> None:
    pf = paragraph.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, fld_end])


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Title", 26, BLACK, 0, 14),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
        ("Heading 4", 11, DEEP_BLUE, 6, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT_BODY)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = FONT_BODY
        code._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        code._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        code._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        code._element.rPr.rFonts.set(qn("w:cs"), FONT_BODY)
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor.from_string("E6FFE6")
        code.paragraph_format.left_indent = Inches(0.16)
        code.paragraph_format.right_indent = Inches(0.16)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(8)
        code.paragraph_format.line_spacing = 1.08

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.columns[0].width = Inches(4.8)
    table.columns[1].width = Inches(1.7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    left.text = "POCKET EARTH · QWEN + MNN"
    right.text = "FINAL · 2026.08.13"
    for i, cell in enumerate((left, right)):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 20, 0, 40, 0)
        set_cell_border(cell, bottom={"val": "single", "sz": "10", "color": BLACK})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            set_run_font(run, 8.5, True, BLACK, "Courier New")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Pocket Earth 决赛完整产品文档  ·  ")
    set_run_font(run, 8.5, False, GRAY)
    add_page_field(p)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(1.1)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("POCKET\nEARTH")
    set_run_font(r, 38, True, BLACK, "Courier New")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("阿里通义 × Arm 手机 AI 挑战赛\n决赛完整产品文档")
    set_run_font(r, 21, True, BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("把现实世界，装进一颗可验证的私人地球")
    set_run_font(r, 15, True, DEEP_BLUE)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(4.85)
    rows = [
        ("MODEL", "Qwen3‑VL‑2B"),
        ("RUNTIME", "MNN 3.6.1 · Android ARM64"),
        ("SYSTEM", "PHOTOS / EARTH / SKILLS · Frost Harness"),
        ("VERSION", "Final 1.0 · 2026-08-13"),
    ]
    for idx, (k, v) in enumerate(rows):
        c0, c1 = table.rows[idx].cells
        c0.text = k
        c1.text = v
        set_cell_shading(c0, BLACK)
        set_cell_shading(c1, PALE_GREEN if idx == 0 else LIGHT)
        for c in (c0, c1):
            set_cell_margins(c, 110, 130, 110, 130)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(c,
                            top={"val": "single", "sz": "8", "color": BLACK},
                            bottom={"val": "single", "sz": "8", "color": BLACK},
                            start={"val": "single", "sz": "8", "color": BLACK},
                            end={"val": "single", "sz": "8", "color": BLACK})
        for run in c0.paragraphs[0].runs:
            set_run_font(run, 9.5, True, GREEN, "Courier New")
        for run in c1.paragraphs[0].runs:
            set_run_font(run, 10.5, idx == 0, BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("CARRY THE COSMOS")
    set_run_font(r, 11, True, GREEN, "Courier New")
    p = doc.add_paragraph()
    r = p.add_run("产品说明 · 技术架构 · 证据索引 · 30 页 PPT 逐页脚本")
    set_run_font(r, 9.5, False, GRAY)

    doc.add_page_break()
    doc.add_heading("目录", level=1)
    toc_items = [
        "0  文档使用说明与事实口径",
        "1  执行摘要",
        "2  用户、场景与核心痛点",
        "3  产品总体结构",
        "4  Frost：真正的轻量 Agent Harness",
        "5  两种 Skill：流程知识与模型权重",
        "6  Skill Protocol、Data Pack 与私人安装",
        "7  端侧技术架构",
        "8  质量门、Trace 与真机验收账本",
        "9  PHOTOS：端侧照片雷达",
        "10  EARTH：从结果页到行动空间",
        "11  SKILLS：私人能力库与 Agent World",
        "12–15  重点业务 Skills",
        "16  隐私、安全与部署",
        "17  竞赛要求对齐",
        "18  决赛现场主 Demo",
        "19  可验证成果与诚实边界",
        "20  产品路线图",
        "附录 A  30 页横版 PPT 逐页脚本",
        "附录 B–C  证据路径索引与提交前终审",
    ]
    table = doc.add_table(rows=10, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, label in enumerate(toc_items):
        row = idx % 10
        col = idx // 10
        cell = table.cell(row, col)
        set_cell_margins(cell, 65, 100, 65, 100)
        set_cell_shading(cell, "FFFFFF" if row % 2 else LIGHT)
        set_cell_border(cell, bottom={"val": "single", "sz": "3", "color": "D6DCE4"})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, label, 9.5, BLACK)
    # The first numbered section inserts its own page break. Avoid a blank page.


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = 6.5
    widths = [usable / cols] * cols
    if cols == 3:
        widths = [1.35, 2.6, 2.55]
    elif cols == 2:
        widths = [2.0, 4.5]
    for c, width in enumerate(widths):
        for cell in table.columns[c].cells:
            cell.width = Inches(width)
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.text = ""
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, LIGHT if ri == 0 else "FFFFFF")
            set_cell_border(cell,
                            top={"val": "single", "sz": "5", "color": "AAB2BD"},
                            bottom={"val": "single", "sz": "5", "color": "AAB2BD"},
                            start={"val": "single", "sz": "5", "color": "AAB2BD"},
                            end={"val": "single", "sz": "5", "color": "AAB2BD"})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, row[ci] if ci < len(row) else "", 8.7, BLACK)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc: Document, text: str, ordered=False, level=0) -> None:
    p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Inches(0.375 + 0.22 * level)
    pf.first_line_indent = Inches(-0.194)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.208
    add_inline(p, text, 10.5, BLACK)


def add_callout(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, CREAM)
    set_cell_margins(cell, 130, 160, 130, 160)
    set_cell_border(cell,
                    start={"val": "single", "sz": "18", "color": GOLD},
                    top={"val": "single", "sz": "4", "color": "D9CFB1"},
                    bottom={"val": "single", "sz": "4", "color": "D9CFB1"},
                    end={"val": "single", "sz": "4", "color": "D9CFB1"})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, " ".join(lines), 10.5, DEEP_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), BLACK)
    pPr.append(shd)
    r = p.add_run(text)
    set_run_font(r, 9, False, "E6FFE6", FONT_BODY)


def should_break_before_heading(level: int, text: str) -> bool:
    if level == 1:
        return True
    if level == 2 and (re.match(r"^\d+\.", text) or re.match(r"^P\d+", text)):
        return True
    return False


def parse_markdown(doc: Document, content: str) -> None:
    lines = content.splitlines()
    # Cover replaces the first title and metadata block.
    start = next((i for i, line in enumerate(lines) if line.startswith("## 0.")), 0)
    i = start
    code_mode = False
    code_lines: list[str] = []
    quote_lines: list[str] = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if code_mode:
            if stripped.startswith("```"):
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                code_mode = False
            else:
                code_lines.append(line)
            i += 1
            continue
        if stripped.startswith("```"):
            code_mode = True
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines.append(stripped[1:].strip())
            i += 1
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            add_callout(doc, quote_lines)
            quote_lines = []
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            rows: list[list[str]] = []
            rows.append([c.strip() for c in stripped.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue

        h = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if h:
            level = len(h.group(1))
            text = h.group(2)
            if should_break_before_heading(level, text):
                doc.add_page_break()
            p = doc.add_heading(text, level=level)
            if re.match(r"^P\d+", text):
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "EAF7EF")
                pPr.append(shd)
            i += 1
            continue

        if stripped in ("---", "***"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:color"), BLACK)
            borders.append(bottom)
            pPr.append(borders)
            i += 1
            continue

        if stripped.startswith("- [ ]"):
            add_bullet(doc, "☐ " + stripped[5:].strip())
            i += 1
            continue
        if stripped.lower().startswith("- [x]"):
            add_bullet(doc, "☒ " + stripped[5:].strip())
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.+)$", stripped)
        if m:
            add_bullet(doc, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if m:
            add_bullet(doc, m.group(2), ordered=True)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Join wrapped prose until the next markdown structure.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                break
            if (nxt.startswith("#") or nxt.startswith(">") or nxt.startswith("```") or
                    nxt.startswith("|") or re.match(r"^[-*]\s+", nxt) or
                    re.match(r"^\d+\.\s+", nxt) or nxt in ("---", "***")):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        set_paragraph_body(p)
        add_inline(p, " ".join(paragraph_lines), 10.5, BLACK)


def update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    update_fields(doc)
    core = doc.core_properties
    core.title = "Pocket Earth｜阿里通义 × Arm 手机 AI 挑战赛决赛完整产品文档"
    core.subject = "产品说明、技术架构、证据索引与 30 页 PPT 逐页脚本"
    core.author = "Pocket Earth"
    core.keywords = "Pocket Earth, Qwen3-VL-2B, MNN, SME2, Skills, Frost"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
