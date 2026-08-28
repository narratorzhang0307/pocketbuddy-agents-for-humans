#!/usr/bin/env python3
"""Build the frozen Pocket Earth Photos plan DOCX from its Markdown source."""

from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path

from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = RGBColor(20, 28, 34)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 105)
GREEN = RGBColor(8, 122, 67)
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "EEF8ED"
LIGHT_GOLD = "FFF8DC"
CONTENT_DXA = 9360


def set_font(run, size: float | None = None, bold: bool | None = None,
             color: RGBColor | None = None, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def paragraph_border(paragraph, color: str = "000000", size: str = "12") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_cell_fill(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), "120")
    tblind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, 8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "POCKET EARTH · PHOTOS FINAL EXECUTION PLAN"
    set_font(header.runs[0], 8, True, MUTED)
    paragraph_border(header, "D7DEE5", "6")
    add_page_number(section.footer.paragraphs[0])


def add_inline(paragraph, text: str, base_size: float = 11, base_color: RGBColor = INK) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor:match.start()]), base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), base_size, True, base_color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, base_size - 0.5, color=DARK_BLUE, name="Menlo")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]), base_size, color=base_color)


def add_masthead(doc: Document, title: str) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    set_font(kicker.add_run("FINAL · EXECUTION / EVIDENCE / RELEASE GATE"), 9, True, GREEN)
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(5)
    set_font(heading.add_run(title), 23, True, INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    set_font(subtitle.add_run("真实相册轻路由、Qwen3-VL 审美策展、个人偏好账本与真机 MNN/SME2 证据闭环"), 12.5, color=MUTED)
    for label, value in (
        ("版本", "2026-08-12 · 最终冻结版"),
        ("目标形态", "Android Capacitor APK · 原片保留在系统相册"),
        ("评测原则", "未见盲集优先 · 换位 A/B · 配对统计 · 失败不冒充发布"),
        ("比赛技术栈", "Qwen3-VL-2B · MNN 3.6.1 · Arm SME2（真机验收）"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}："), 10.5, True, INK)
        set_font(p.add_run(value), 10.5, color=MUTED)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(14)
    paragraph_border(rule, "111111", "16")


def split_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    widths = [CONTENT_DXA // columns] * columns
    widths[-1] += CONTENT_DXA - sum(widths)
    table = doc.add_table(rows=len(rows), cols=columns)
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            if row_index == 0:
                set_cell_fill(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            text = values[column_index] if column_index < len(values) else ""
            add_inline(paragraph, text, 9.5, INK)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.2
    shade_paragraph(paragraph, LIGHT_GREEN)
    add_inline(paragraph, text, 11, DARK_BLUE)


def render_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.space_after = Pt(8)
                shade_paragraph(paragraph, "F2F4F7")
                set_font(paragraph.add_run("\n".join(code_lines)), 8.5, color=INK, name="Menlo")
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_paragraph(stripped[3:], style="Heading 1")
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:], style="Heading 2")
            index += 1
            continue
        if stripped.startswith("#### "):
            doc.add_paragraph(stripped[5:], style="Heading 3")
            index += 1
            continue
        if stripped.startswith("> "):
            add_quote(doc, stripped[2:])
            index += 1
            continue
        if "|" in stripped and index + 1 < len(lines) and is_table_separator(lines[index + 1].strip()):
            rows = [split_cells(stripped)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(split_cells(lines[index]))
                index += 1
            add_table(doc, rows)
            continue
        if re.match(r"^[-*] ", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            # Preserve the source number explicitly. Word reuses a single
            # automatic numbering instance across separated lists, which can
            # otherwise turn a new 1–5 section into 7–11 after rendering.
            match = re.match(r"^(\d+)\. (.*)$", stripped)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.375)
            paragraph.paragraph_format.first_line_indent = Inches(-0.188)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            set_font(paragraph.add_run(f"{match.group(1)}.  "), 11, color=INK)
            add_inline(paragraph, match.group(2))
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", ">", "```")) or re.match(r"^[-*] |^\d+\. ", candidate):
                break
            if "|" in candidate and index + 1 < len(lines) and is_table_separator(lines[index + 1].strip()):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def add_evidence_gallery(doc: Document, evidence_dir: Path) -> None:
    images = [
        ("00-task-pack-overview.png", "外置任务包：32 张照片、EXIF/GPS 与预期结果"),
        ("02-curation-pipeline-and-groups.png", "真实轻索引后的决策矩阵与分组"),
        ("03-burst-human-decision.png", "连拍组：技术门、Qwen 策展与用户最终确认分离"),
        ("04-on-device-proof-web-truthful.png", "证据面板：网页预览明确不冒充真机 MNN/SME2"),
    ]
    available = [(evidence_dir / name, caption) for name, caption in images if (evidence_dir / name).is_file()]
    if not available:
        return
    doc.add_page_break()
    doc.add_paragraph("PPT 证据页", style="Heading 1")
    intro = doc.add_paragraph()
    add_inline(intro, "以下均为本轮真实任务包和真实产品链路截图；可直接裁入答辩 PPT。", 10.5, MUTED)
    with tempfile.TemporaryDirectory(prefix="pocket-earth-docx-evidence-") as temp_dir:
        for position, (path, caption) in enumerate(available):
            picture_path = path
            if path.name != "00-task-pack-overview.png":
                with Image.open(path).convert("RGB") as source:
                    background = Image.new("RGB", source.size, source.getpixel((0, 0)))
                    difference = ImageChops.difference(source, background).convert("L").point(lambda value: 255 if value > 8 else 0)
                    bbox = difference.getbbox()
                    if bbox:
                        left, top, right, bottom = bbox
                        padding = 12
                        crop_box = (
                            max(0, left - padding), max(0, top - padding),
                            min(source.width, right + padding), min(source.height, bottom + padding),
                        )
                        picture_path = Path(temp_dir) / path.name
                        source.crop(crop_box).save(picture_path)
            picture_paragraph = doc.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.keep_with_next = True
            if position:
                picture_paragraph.paragraph_format.page_break_before = True
            if path.name == "00-task-pack-overview.png":
                picture_paragraph.add_run().add_picture(str(picture_path), height=Inches(6.25))
            else:
                picture_paragraph.add_run().add_picture(str(picture_path), height=Inches(7.15))
            note = doc.add_paragraph()
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            note.paragraph_format.space_before = Pt(6)
            set_font(note.add_run(caption), 9.5, True, DARK_BLUE)


def build(markdown_path: Path, output_path: Path, evidence_dir: Path | None) -> None:
    markdown = markdown_path.read_text(encoding="utf-8")
    title = next((line[2:].strip() for line in markdown.splitlines() if line.startswith("# ")), "Pocket Earth Photos 最终执行计划")
    doc = Document()
    configure_document(doc)
    add_masthead(doc, title)
    render_markdown(doc, markdown)
    if evidence_dir:
        add_evidence_gallery(doc, evidence_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve(), args.evidence_dir.resolve() if args.evidence_dir else None)
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
